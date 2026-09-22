"""Test cho service import sinh viên từ Excel.

Bao gồm các trường hợp:
    - File hợp lệ: tất cả dòng được tạo, mật khẩu là ddmmyyyy từ ngày sinh.
    - Một dòng lỗi không rollback các dòng đã tạo trước đó (per-row commit).
    - Trùng user_id/email trong file và trong DB.
    - File sai định dạng / thiếu cột / vượt quá giới hạn / trống.
"""

from datetime import date

import pytest
from openpyxl import Workbook
from sqlmodel import select

from app.core.password import verify_password
from app.services.student_service import (
    MAX_IMPORT_ROWS,
    EmptyExcelFileError,
    ExcelFileTooLargeError,
    InvalidExcelFormatError,
    import_students,
)
from models import User


def _build_xlsx(
    rows: list[tuple[str, str, str, str]],
    header: tuple[str, ...] = ("user_id", "full_name", "date_of_birth", "email"),
) -> bytes:
    """Tạo file .xlsx trong bộ nhớ từ danh sách dòng dữ liệu."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    from io import BytesIO

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_import_students_success_creates_accounts_with_date_based_password(session):
    contents = _build_xlsx(
        [
            ("sv001", "Nguyen Van A", "31/08/2005", "sv001@st.utc2.edu.vn"),
            ("sv002", "Nguyen Van B", "01/01/2004", "sv002@st.utc2.edu.vn"),
        ]
    )

    response = import_students(session, contents, "students.xlsx")

    assert len(response.created) == 2
    assert response.errors == []

    # Mật khẩu tạm thời phải là ddmmyyyy từ ngày sinh
    passwords = {row.user_id: row.temporary_password for row in response.created}
    assert passwords["sv001"] == "31082005"
    assert passwords["sv002"] == "01012004"

    users = session.exec(select(User)).all()
    assert {u.user_id for u in users} == {"sv001", "sv002"}
    for user in users:
        assert verify_password(
            passwords[user.user_id], user.hashed_password
        )
        assert user.must_change_password is True
        assert user.role_id == 3


def test_import_students_accepts_iso_date_of_birth(session):
    contents = _build_xlsx(
        [("sv003", "Nguyen Van C", "2005-08-31", "sv003@st.utc2.edu.vn")]
    )

    response = import_students(session, contents, "students.xlsx")

    assert len(response.created) == 1
    assert response.created[0].temporary_password == "31082005"


def _build_xlsx_with_typed_cells(
    rows: list[tuple[object, object, object, object]],
    header: tuple[str, ...] = ("user_id", "full_name", "date_of_birth", "email"),
) -> bytes:
    """Tạo file .xlsx cho phép truyền cell với kiểu dữ liệu tuỳ ý."""
    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    from io import BytesIO

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_import_students_accepts_excel_date_cell_format(session):
    """Khi cột date_of_birth được định dạng kiểu Date trong Excel, openpyxl trả về
    ``datetime.date`` thay vì chuỗi — service vẫn phải parse đúng và sinh
    mật khẩu ddmmyyyy."""
    from datetime import date as _date

    contents = _build_xlsx_with_typed_cells(
        [
            ("sv_dt", "Nguyen Van Date", _date(2005, 8, 31), "sv_dt@st.utc2.edu.vn"),
        ]
    )

    response = import_students(session, contents, "students.xlsx")

    assert response.errors == []
    assert len(response.created) == 1
    assert response.created[0].temporary_password == "31082005"


def test_import_students_accepts_excel_datetime_cell_format(session):
    """Khi cột date_of_birth là ô ``Date`` có giờ phút giây, openpyxl trả về
    ``datetime.datetime``. Nếu chỉ ``str()`` thì ra ``"yyyy-mm-dd hh:mm:ss"``
    và bị ``parse_date_of_birth`` reject — service phải convert về ISO date."""
    from datetime import datetime as _datetime

    contents = _build_xlsx_with_typed_cells(
        [
            (
                "sv_dt2",
                "Nguyen Van Datetime",
                _datetime(2005, 8, 31, 14, 30, 0),
                "sv_dt2@st.utc2.edu.vn",
            ),
        ]
    )

    response = import_students(session, contents, "students.xlsx")

    assert response.errors == []
    assert len(response.created) == 1
    assert response.created[0].temporary_password == "31082005"


def test_import_students_does_not_block_other_rows_when_one_fails(session):
    """Một dòng lỗi không được rollback các dòng đã tạo thành công trước đó."""
    contents = _build_xlsx(
        [
            ("sv_ok_1", "Nguyen Van A", "31/08/2005", "sv_ok_1@st.utc2.edu.vn"),
            (
                "sv_dup",
                "Nguyen Van B",
                "31/08/2005",
                "sv_ok_1@st.utc2.edu.vn",
            ),  # trùng email
            ("sv_ok_2", "Nguyen Van C", "01/01/2004", "sv_ok_2@st.utc2.edu.vn"),
        ]
    )

    response = import_students(session, contents, "students.xlsx")

    created_ids = {row.user_id for row in response.created}
    assert created_ids == {"sv_ok_1", "sv_ok_2"}
    assert len(response.errors) == 1
    assert response.errors[0].row == 3

    # Cả 2 dòng thành công đều đã có mặt trong DB
    users = session.exec(select(User)).all()
    assert {u.user_id for u in users} == {"sv_ok_1", "sv_ok_2"}


def test_import_students_rejects_invalid_excel_format(session):
    with pytest.raises(InvalidExcelFormatError):
        import_students(session, b"abc", "students.csv")


def test_import_students_rejects_empty_file(session):
    contents = _build_xlsx([])  # chỉ có header

    with pytest.raises(EmptyExcelFileError):
        import_students(session, contents, "students.xlsx")


def test_import_students_rejects_missing_required_columns(session):
    wb = Workbook()
    ws = wb.active
    ws.append(["full_name", "email"])  # thiếu user_id, date_of_birth
    ws.append(["A", "a@example.com"])
    from io import BytesIO

    buffer = BytesIO()
    wb.save(buffer)

    with pytest.raises(InvalidExcelFormatError):
        import_students(session, buffer.getvalue(), "students.xlsx")


def test_import_students_rejects_invalid_date_of_birth(session):
    """Ngày không đúng 3 format dd/MM/yyyy, dd-MM-yyyy, yyyy-mm-dd bị reject."""
    contents = _build_xlsx(
        [("sv004", "Nguyen Van D", "31/08/05", "sv004@st.utc2.edu.vn")]
    )

    response = import_students(session, contents, "students.xlsx")
    assert response.created == []
    assert len(response.errors) == 1
    assert "định dạng" in response.errors[0].reason.lower()


def test_import_students_rejects_too_many_rows(session):
    rows = [
        (f"sv{i:04d}", f"Nguyen Van {i}", "31/08/2005", f"sv{i:04d}@st.utc2.edu.vn")
        for i in range(MAX_IMPORT_ROWS + 1)
    ]
    contents = _build_xlsx(rows)

    with pytest.raises(ExcelFileTooLargeError):
        import_students(session, contents, "students.xlsx")


def test_import_students_rejects_duplicate_with_existing_db(session):
    existing = User(
        user_id="existing",
        full_name="Existing",
        email="existing@st.utc2.edu.vn",
        hashed_password="placeholder",
        role_id=3,
        date_of_birth=date(2005, 8, 31),
    )
    session.add(existing)
    session.commit()

    contents = _build_xlsx(
        [
            ("existing", "Trung Lap", "31/08/2005", "new_email@st.utc2.edu.vn"),
        ]
    )

    response = import_students(session, contents, "students.xlsx")

    assert response.created == []
    assert len(response.errors) == 1
    assert "đã tồn tại" in response.errors[0].reason.lower()
