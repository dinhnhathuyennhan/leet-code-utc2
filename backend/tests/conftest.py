import os
import sys
import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import Workbook
from sqlmodel import Session, SQLModel, create_engine
from sqlmodel.pool import StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import models  # noqa: F401  # phải sau sys.path.insert để tìm được module 'models' ở backend/

# Load .env ở root project TRƯỚC khi import app modules
# (cần cho JWT_SECRET_KEY — create_access_token require string, không được None).
_root_env = Path(__file__).resolve().parents[2] / ".env"
if _root_env.exists():
    from dotenv import load_dotenv

    load_dotenv(_root_env, override=False)

# Fallback nếu vẫn None (an toàn cho test).
if not os.getenv("JWT_SECRET_KEY"):
    os.environ["JWT_SECRET_KEY"] = "test-secret-key-for-pytest"

# Phần đầu sheet giống file "DS lớp học phần" thật của phòng đào tạo (hàng 1-10)
CLASS_LIST_PREAMBLE = [
    ["TRƯỜNG ĐH GIAO THÔNG VẬN TẢI", None, None, None, None, None, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM"],
    ["PHÂN HIỆU TẠI TP. HỒ CHÍ MINH", None, None, None, None, None, "Độc lập - Tự do - Hạnh phúc"],
    [],
    [".", ".", ".", ".", "."],
    ["DANH SÁCH ĐIỂM DANH SINH VIÊN"],
    ["Lớp học phần:", None, "Nhập môn ngành CNTT-1-26-CQ.67.CNTT"],
    ["Mã học phần:", None, "Nhập môn ngành CNTT-IT1.001.3"],
    ["Học kỳ: ", None, "Học kỳ: 1 - Năm học: 2026_2027"],
    ["Thời khóa biểu:", None, "Thời gian học: 14/09/2026->14/11/2026;GV: Giảng viên A"],
    [],
]
CLASS_LIST_HEADER = ["STT", "Mã sinh viên", "Họ và tên", None, "Lớp", None, "Ghi chú"]


def build_class_list_xlsx(sheets: dict[str, list[tuple] | None]) -> bytes:
    """Dựng file .xlsx theo đúng bố cục file "DS lớp học phần" thật.

    ``sheets``: tên sheet → danh sách sinh viên ``(mssv, họ đệm, tên, lớp)``.
    Phần tử ``None`` trong danh sách là một dòng trống hoàn toàn (không có
    STT). Giá trị ``None`` thay cho cả danh sách nghĩa là sheet không theo
    mẫu (không có hàng tiêu đề). Mỗi sheet gồm phần đầu 10 hàng, tiêu đề ở
    hàng 11 (ô "Họ và tên" gộp 2 cột), dữ liệu từ hàng 12, rồi chân trang.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, students in sheets.items():
        worksheet = workbook.create_sheet(title)
        if students is None:
            worksheet.append(["Ghi chú", "Sheet này không phải danh sách lớp"])
            continue
        for line in CLASS_LIST_PREAMBLE:
            worksheet.append(line)
        worksheet.append(CLASS_LIST_HEADER)
        worksheet.merge_cells("C11:D11")
        stt = 0
        for student in students:
            if student is None:
                worksheet.append([])
                continue
            stt += 1
            worksheet.append([stt, *student])
        worksheet.append([])
        worksheet.append([None, None, None, None, None, None, "TP Hồ Chí Minh, ngày 5 tháng 10 năm 2026"])
        worksheet.append(["Ghi chú: giáo viên điểm danh ngày nào thì ghi ngày điểm danh vào đầu cột"])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture()
def broken_sheet_xlsx() -> bytes:
    """File .xlsx mở được (zip và workbook hợp lệ) nhưng XML của sheet bị cắt cụt.

    openpyxl ở chế độ ``read_only`` chỉ đọc XML của sheet khi duyệt hàng,
    nên lỗi chỉ lộ ra lúc đọc ô chứ không phải lúc mở file.
    """
    workbook = Workbook()
    workbook.active.append(["STT", "Mã sinh viên"])
    valid = BytesIO()
    workbook.save(valid)

    broken = BytesIO()
    with zipfile.ZipFile(BytesIO(valid.getvalue())) as source, zipfile.ZipFile(broken, "w") as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = b"<worksheet><sheetData><row><c r="
            target.writestr(name, data)
    return broken.getvalue()


@pytest.fixture()
def class_list_xlsx():
    """Trả về hàm ``build_class_list_xlsx`` để test tự dựng file Excel."""
    return build_class_list_xlsx


# Tách engine và session ra riêng để dùng linh động hơn trong việc test
@pytest.fixture()
def engine():
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    SQLModel.metadata.create_all(engine)
    return engine


@pytest.fixture()
def session(engine):
    with Session(engine) as s:
        yield s
