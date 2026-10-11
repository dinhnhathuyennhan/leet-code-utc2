"""Test module đọc Excel → JSON (``app.core.excel.read_workbook``).

File ``fixtures/ds_lop_mau.xls`` (dữ liệu giả) được tạo bằng cách dựng file
.xlsx với ``build_class_list_xlsx`` (3 sinh viên, MSSV đầu tiên là ô số) cộng
thêm sheet "Ghi chu" có ô ngày 05/10/2026, rồi chuyển sang .xls bằng
``soffice --headless --convert-to xls``.
"""

import zipfile
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import pytest
import xlrd
from openpyxl import Workbook, load_workbook

from app.core.excel import (
    EmptyExcelFileError,
    ExcelFileTooLargeError,
    InvalidExcelFormatError,
    read_workbook,
)
MAX_BYTES = 5 * 1024 * 1024
XLS_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "ds_lop_mau.xls"


def _xlsx(*sheets: tuple[str, list[list]]) -> bytes:
    """Dựng file .xlsx từ các cặp ``(tên sheet, danh sách hàng)``."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets:
        worksheet = workbook.create_sheet(title)
        for row in rows:
            worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ============================== ĐỌC .xlsx ==============================


def test_read_xlsx_converts_every_cell_to_string():
    contents = _xlsx(
        (
            "S1",
            [
                ["STT", "Mã sinh viên", "Điểm", "Ngày", "Giờ"],
                [1, 6751071001, 3.5, date(2005, 8, 31), datetime(2005, 8, 31, 7, 30)],
                ["  2 ", " 675107A002 ", None, "31/08/2005", True],
            ],
        )
    )

    workbook = read_workbook(contents, "ds.xlsx", MAX_BYTES)

    assert workbook.model_dump() == {
        "sheets": [
            {
                "name": "S1",
                "rows": [
                    ["STT", "Mã sinh viên", "Điểm", "Ngày", "Giờ"],
                    ["1", "6751071001", "3.5", "2005-08-31", "2005-08-31"],
                    ["2", "675107A002", "", "31/08/2005", "True"],
                ],
            }
        ]
    }


def test_read_xlsx_trims_trailing_empty_cells_and_rows():
    contents = _xlsx(
        (
            "S1",
            [
                ["a", None, "b", None, None],
                [],
                ["c"],
                [None, None],
                [],
            ],
        )
    )

    rows = read_workbook(contents, "ds.xlsx", MAX_BYTES).sheets[0].rows

    # Ô trống ở giữa và hàng trống ở giữa vẫn giữ để chỉ số hàng khớp với Excel
    assert rows == [["a", "", "b"], [], ["c"]]


def test_read_xlsx_reads_all_sheets_and_drops_empty_sheets():
    contents = _xlsx(
        ("Lop1", [["x"]]),
        ("Trong", []),
        ("Lop2", [["y"]]),
    )

    workbook = read_workbook(contents, "ds.xlsx", MAX_BYTES)

    assert [sheet.name for sheet in workbook.sheets] == ["Lop1", "Lop2"]


def test_read_xlsx_accepts_uppercase_extension():
    contents = _xlsx(("S1", [["x"]]))

    workbook = read_workbook(contents, "DS LOP.XLSX", MAX_BYTES)

    assert workbook.sheets[0].rows == [["x"]]


def test_read_xlsx_class_list_layout(class_list_xlsx):
    contents = class_list_xlsx(
        {"IT1.001": [(6751071001, "Nguyễn Văn", "An", "CQ.67.CNTT")]}
    )

    rows = read_workbook(contents, "ds.xlsx", MAX_BYTES).sheets[0].rows

    assert rows[10] == ["STT", "Mã sinh viên", "Họ và tên", "", "Lớp", "", "Ghi chú"]
    assert rows[11] == ["1", "6751071001", "Nguyễn Văn", "An", "CQ.67.CNTT"]


# ============================== ĐỌC .xls ==============================


def test_read_xls_fixture():
    workbook = read_workbook(XLS_FIXTURE.read_bytes(), "ds_lop_mau.xls", MAX_BYTES)

    assert [sheet.name for sheet in workbook.sheets] == [
        "IT1.001.3-1-26-CQ.67.CNTT",
        "Ghi chu",
    ]
    rows = workbook.sheets[0].rows
    assert rows[10] == ["STT", "Mã sinh viên", "Họ và tên", "", "Lớp", "", "Ghi chú"]
    # MSSV lưu dạng ô số trong .xls (6751071001.0) phải ra chuỗi không có ".0"
    assert rows[11] == ["1", "6751071001", "Nguyễn Văn", "An", "CQ.67.CNTT"]
    assert rows[12] == ["2", "675107A002", "Trần Thị", "Bình", "CQ.67.CNTT"]
    assert rows[13] == ["3", "A6751071003", "Lê", "Cường", "CQ.67.KTXDCTGT.2"]
    # Ô ngày trong .xls (lưu dạng số thực) đổi về yyyy-mm-dd
    assert workbook.sheets[1].rows == [["Ngày lập", "2026-10-05"]]


def test_xls_and_xlsx_of_same_content_give_same_json(class_list_xlsx):
    """Fixture .xls được chuyển từ đúng file .xlsx này nên JSON phải giống hệt."""
    xlsx = _xlsx_like_fixture(class_list_xlsx)

    from_xlsx = read_workbook(xlsx, "ds.xlsx", MAX_BYTES)
    from_xls = read_workbook(XLS_FIXTURE.read_bytes(), "ds.xls", MAX_BYTES)

    assert from_xls.model_dump() == from_xlsx.model_dump()


def _xlsx_like_fixture(class_list_xlsx) -> bytes:
    contents = class_list_xlsx(
        {
            "IT1.001.3-1-26-CQ.67.CNTT": [
                (6751071001, "Nguyễn Văn", "An", "CQ.67.CNTT"),
                ("675107A002", "Trần Thị", "Bình", "CQ.67.CNTT"),
                ("A6751071003", "Lê", "Cường", "CQ.67.KTXDCTGT.2"),
            ]
        }
    )
    workbook = load_workbook(BytesIO(contents))
    workbook.create_sheet("Ghi chu").append(["Ngày lập", datetime(2026, 10, 5)])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ============================== LỖI ==============================


@pytest.mark.parametrize("filename", ["ds.csv", "ds.xlsm", "ds", "", None])
def test_rejects_unsupported_extension(filename):
    with pytest.raises(InvalidExcelFormatError):
        read_workbook(b"abc", filename, MAX_BYTES)


def test_rejects_corrupt_xlsx():
    with pytest.raises(InvalidExcelFormatError) as exc_info:
        read_workbook(b"not a zip file", "ds.xlsx", MAX_BYTES)

    # Không lộ chi tiết lỗi nội bộ của thư viện ra thông báo
    assert exc_info.value.message == "Không đọc được file Excel"


def test_rejects_zip_that_is_not_excel():
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "hi")

    with pytest.raises(InvalidExcelFormatError):
        read_workbook(buffer.getvalue(), "ds.xlsx", MAX_BYTES)


def test_rejects_corrupt_xls():
    with pytest.raises(InvalidExcelFormatError) as exc_info:
        read_workbook(b"garbage" * 100, "ds.xls", MAX_BYTES)

    assert exc_info.value.message == "Không đọc được file Excel"


def test_rejects_xlsx_with_broken_sheet_data(broken_sheet_xlsx):
    """Lỗi chỉ lộ ra lúc đọc từng hàng (không phải lúc mở file) vẫn phải ra 422."""
    with pytest.raises(InvalidExcelFormatError) as exc_info:
        read_workbook(broken_sheet_xlsx, "ds.xlsx", MAX_BYTES)

    assert exc_info.value.message == "Không đọc được file Excel"


def test_rejects_xls_with_unreadable_date_cell(monkeypatch):
    """Ô ngày có giá trị ngoài phạm vi làm xlrd lỗi khi đổi sang datetime."""

    def broken_date(*_args):
        raise OverflowError("date value out of range")

    monkeypatch.setattr(xlrd, "xldate_as_datetime", broken_date)

    with pytest.raises(InvalidExcelFormatError):
        read_workbook(XLS_FIXTURE.read_bytes(), "ds.xls", MAX_BYTES)


def test_rejects_xlsx_renamed_to_xls():
    contents = _xlsx(("S1", [["x"]]))

    with pytest.raises(InvalidExcelFormatError):
        read_workbook(contents, "ds.xls", MAX_BYTES)


def test_rejects_xls_renamed_to_xlsx():
    with pytest.raises(InvalidExcelFormatError):
        read_workbook(XLS_FIXTURE.read_bytes(), "ds.xlsx", MAX_BYTES)


def test_rejects_file_larger_than_max_bytes():
    contents = _xlsx(("S1", [["x"]]))

    with pytest.raises(ExcelFileTooLargeError):
        read_workbook(contents, "ds.xlsx", len(contents) - 1)


def test_accepts_file_exactly_max_bytes():
    contents = _xlsx(("S1", [["x"]]))

    workbook = read_workbook(contents, "ds.xlsx", len(contents))

    assert workbook.sheets[0].rows == [["x"]]


def test_rejects_workbook_without_any_data():
    contents = _xlsx(("S1", []), ("S2", [[None, None]]))

    with pytest.raises(EmptyExcelFileError):
        read_workbook(contents, "ds.xlsx", MAX_BYTES)
