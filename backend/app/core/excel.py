"""Đọc file Excel (.xls / .xlsx) thành dữ liệu dạng JSON trong bộ nhớ.

Module này không biết gì về nghiệp vụ (sinh viên, lớp...). Nó chỉ chuyển
cả workbook thành ``ExcelWorkbook``: danh sách sheet, mỗi sheet là danh
sách hàng, mỗi hàng là danh sách ô đã chuẩn hoá thành chuỗi. Service gọi
module này rồi tự bốc tách dữ liệu theo mẫu file của mình.
"""

from datetime import date as _date
from datetime import datetime as _datetime
from io import BytesIO
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel

from app.core.exceptions import AppError


class InvalidExcelFormatError(AppError):
    status_code = 422
    error_code = "INVALID_EXCEL_FORMAT"


class EmptyExcelFileError(AppError):
    status_code = 400
    error_code = "EMPTY_EXCEL_FILE"


class ExcelFileTooLargeError(AppError):
    status_code = 400
    error_code = "EXCEL_FILE_TOO_LARGE"


class ExcelSheet(BaseModel):
    """Một sheet. ``rows[i]`` là hàng thứ ``i + 1`` trong Excel."""

    name: str
    rows: list[list[str]]


class ExcelWorkbook(BaseModel):
    """Toàn bộ file Excel ở dạng JSON (gọi ``.model_dump()`` để xem)."""

    sheets: list[ExcelSheet]


def read_workbook(contents: bytes, filename: str, max_bytes: int) -> ExcelWorkbook:
    """Đọc file .xls hoặc .xlsx thành ``ExcelWorkbook``.

    - Mọi ô đều là chuỗi đã strip; ô trống là ``""``.
    - Ô số nguyên (vd MSSV lưu dạng số ``6751071095.0``) thành ``"6751071095"``.
    - Ô ngày thành ``"yyyy-mm-dd"``.
    - Bỏ các ô trống ở cuối mỗi hàng và các hàng trống ở cuối sheet,
      nên các hàng có thể dài ngắn khác nhau. Sheet không có dữ liệu bị bỏ qua.

    Raises:
        InvalidExcelFormatError, ExcelFileTooLargeError, EmptyExcelFileError.
    """
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        reader = _read_xlsx
    elif name.endswith(".xls"):
        reader = _read_xls
    else:
        raise InvalidExcelFormatError("File phải có định dạng .xls hoặc .xlsx")

    if len(contents) > max_bytes:
        raise ExcelFileTooLargeError(
            f"File vượt quá dung lượng {max_bytes // (1024 * 1024)} MB cho phép"
        )

    sheets = [sheet for sheet in reader(contents) if sheet.rows]
    if not sheets:
        raise EmptyExcelFileError("File không có dữ liệu")
    return ExcelWorkbook(sheets=sheets)


def _read_xlsx(contents: bytes) -> list[ExcelSheet]:
    try:
        workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError) as err:
        raise InvalidExcelFormatError("Không đọc được file Excel") from err

    try:
        return [
            ExcelSheet(
                name=worksheet.title,
                rows=_clean_rows(
                    [_cell_to_string(value) for value in row]
                    for row in worksheet.iter_rows(values_only=True)
                ),
            )
            for worksheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _read_xls(contents: bytes) -> list[ExcelSheet]:
    try:
        book = xlrd.open_workbook(file_contents=contents)
    except Exception as err:  # xlrd raise nhiều loại lỗi khác nhau khi file hỏng
        raise InvalidExcelFormatError("Không đọc được file Excel") from err

    sheets: list[ExcelSheet] = []
    for sheet in book.sheets():
        rows = (
            [_xls_cell_to_string(cell, book.datemode) for cell in sheet.row(index)]
            for index in range(sheet.nrows)
        )
        sheets.append(ExcelSheet(name=sheet.name, rows=_clean_rows(rows)))
    return sheets


def _xls_cell_to_string(cell: xlrd.sheet.Cell, datemode: int) -> str:
    """xlrd lưu ngày dạng số thực, phải đổi sang datetime trước."""
    if cell.ctype == xlrd.XL_CELL_DATE:
        return _cell_to_string(xlrd.xldate_as_datetime(cell.value, datemode))
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return ""
    return _cell_to_string(cell.value)


def _cell_to_string(value) -> str:
    """Chuẩn hoá giá trị một ô thành chuỗi."""
    if value is None:
        return ""
    if isinstance(value, _datetime):
        return value.date().isoformat()
    if isinstance(value, _date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_rows(rows) -> list[list[str]]:
    """Bỏ ô trống cuối mỗi hàng, rồi bỏ các hàng trống ở cuối sheet."""
    cleaned: list[list[str]] = []
    for row in rows:
        while row and row[-1] == "":
            row.pop()
        cleaned.append(row)
    while cleaned and not cleaned[-1]:
        cleaned.pop()
    return cleaned
